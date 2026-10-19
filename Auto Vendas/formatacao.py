"""criação e formatação da planilha de vendas e do dashboard."""
import math
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

NOME_ABA = "Vendas"
NOME_DASHBOARD = "Dashboard"

COR_CABECALHO = "1F4E78"
COR_BORDA = "D9D9D9"
ALTURA_CABECALHO = 24
LARGURA_MAXIMA = 60

FORMATO_MOEDA = '"R$" #,##0.00'
FORMATO_DATA = "dd/mm/yyyy"
COLUNAS_MONETARIAS = {"Preço unit.", "Total"}

# Alimenta os gráficos.
COLUNA_AUXILIAR = 20

ESTILO_COLUNAS = {
    "ID": ("center", "0"),
    "Pedido": ("center", "@"),
    "Data": ("center", FORMATO_DATA),
    "Cliente": ("left", "General"),
    "Produto": ("left", "General"),
    "Qtd.": ("right", "0"),
    "Preço unit.": ("right", FORMATO_MOEDA),
    "Total": ("right", FORMATO_MOEDA),
}


class ErroDeFormatacao(Exception):
    """Falha ao formatar a planilha de saída."""


def _eh_numero(valor) -> bool:
    return isinstance(valor, (int, float)) and not isinstance(valor, bool)


def _tamanho_exibido(valor, cabecalho=None) -> int:
    """Estima quantos caracteres o valor ocupa na tela."""
    if valor is None:
        return 0
    if isinstance(valor, (datetime, date)):
        return len(FORMATO_DATA)
    if cabecalho in COLUNAS_MONETARIAS and _eh_numero(valor):
        return len(f"R$ {valor:,.2f}")
    return len(str(valor))


def _formatar_cabecalho(ws, borda) -> None:
    preenchimento = PatternFill(
        "solid", start_color=COR_CABECALHO, end_color=COR_CABECALHO
    )
    fonte = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    alinhamento = Alignment(horizontal="center", vertical="center")

    for celula in ws[1]:
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = alinhamento
        celula.border = borda

    ws.row_dimensions[1].height = ALTURA_CABECALHO


def _formatar_corpo(ws, borda) -> None:
    cabecalhos = {celula.column: celula.value for celula in ws[1]}

    for linha in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for celula in linha:
            horizontal, formato = ESTILO_COLUNAS.get(
                cabecalhos.get(celula.column), ("left", "General")
            )
            celula.alignment = Alignment(horizontal=horizontal, vertical="center")
            celula.number_format = formato
            celula.border = borda


def _ajustar_larguras(ws) -> None:
    for coluna in ws.iter_cols(min_row=1, max_row=ws.max_row):
        cabecalho = coluna[0].value
        largura_cabecalho = _tamanho_exibido(cabecalho) + 4  # espaço do filtro
        largura_conteudo = max(
            (_tamanho_exibido(c.value, cabecalho) for c in coluna[1:]),
            default=0,
        ) + 2
        letra = get_column_letter(coluna[0].column)
        ws.column_dimensions[letra].width = min(
            max(largura_cabecalho, largura_conteudo), LARGURA_MAXIMA
        )


def _agrupar_vendas(ws):
    """Agrupa faturamento e quantidade por produto e faturamento por data."""
    cabecalhos = {celula.value: celula.column for celula in ws[1]}
    obrigatorias = {"Produto", "Qtd.", "Total", "Data"}

    if not obrigatorias.issubset(cabecalhos):
        ausentes = sorted(obrigatorias - set(cabecalhos))
        raise ErroDeFormatacao(
            "Coluna(s) necessária(s) para os gráficos ausente(s): "
            f"{', '.join(ausentes)}."
        )

    faturamento_produto = {}
    quantidade_produto = {}
    faturamento_data = {}

    for linha in range(2, ws.max_row + 1):
        produto = ws.cell(linha, cabecalhos["Produto"]).value
        quantidade = ws.cell(linha, cabecalhos["Qtd."]).value
        total = ws.cell(linha, cabecalhos["Total"]).value
        data_venda = ws.cell(linha, cabecalhos["Data"]).value

        if produto is None or not str(produto).strip():
            continue
        produto = str(produto).strip()

        if _eh_numero(total):
            faturamento_produto[produto] = faturamento_produto.get(produto, 0) + total

        if _eh_numero(quantidade):
            quantidade_produto[produto] = quantidade_produto.get(produto, 0) + quantidade

        if _eh_numero(total) and isinstance(data_venda, (date, datetime)):
            dia = data_venda.date() if isinstance(data_venda, datetime) else data_venda
            faturamento_data[dia] = faturamento_data.get(dia, 0) + total

    return faturamento_produto, quantidade_produto, faturamento_data


def _passo_escala(maximo, divisoes, inteiro=False):
    """Escolhe um passo "redondo" (1, 2, 5 x 10^n) para o eixo de valores."""
    bruto = maximo / divisoes
    magnitude = 10 ** math.floor(math.log10(bruto))
    passo = next(f * magnitude for f in (1, 2, 5, 10) if f * magnitude >= bruto)
    return max(1, round(passo)) if inteiro else passo


def _configurar_escala(eixo, maximo, divisoes, inteiro=False) -> None:
    """Fixa o eixo de valores em 0..máximo, com divisões legíveis."""
    if not maximo or maximo <= 0:
        return
    passo = _passo_escala(maximo, divisoes, inteiro)
    eixo.scaling.min = 0
    eixo.scaling.max = round(math.ceil(maximo / passo) * passo, 10)
    eixo.majorUnit = passo


def _preparar_grafico(grafico) -> None:
    """Configura a aparência e os eixos do gráfico."""
    grafico.visible_cells_only = False
    grafico.roundedCorners = False

    if hasattr(grafico, "x_axis"):
        grafico.x_axis.delete = False
        grafico.y_axis.delete = False

    # Contorno do gráfico (atributo c
    grafico.graphical_properties = GraphicalProperties(
        solidFill="FFFFFF",
        ln=LineProperties(solidFill="B7B7B7", w=12700),
    )


def _criar_grafico_barras(
    ws, titulo, coluna_dados, coluna_categorias,
    ultima_linha, posicao, tipo="bar", titulo_eixo="Valor",
    maximo=None, inteiro=False, altura=9,
) -> None:
    """Cria gráficos de barras horizontais ('bar') ou verticais ('col')."""
    grafico = BarChart()
    grafico.type = tipo
    grafico.style = 10
    grafico.title = titulo
    grafico.height = altura
    grafico.width = 17

    dados = Reference(
        ws, min_col=coluna_dados, min_row=1, max_row=ultima_linha
    )
    categorias = Reference(
        ws, min_col=coluna_categorias, min_row=2, max_row=ultima_linha
    )

    grafico.add_data(dados, titles_from_data=True)
    grafico.set_categories(categorias)
    grafico.legend = None

    grafico.y_axis.title = titulo_eixo
    if tipo == "bar":
        grafico.x_axis.axPos = "l"
        grafico.y_axis.axPos = "b"
        grafico.x_axis.scaling.orientation = "maxMin"  #
        grafico.y_axis.crosses = "max"  # números do eixo 
    else:
        grafico.x_axis.axPos = "b"
        grafico.y_axis.axPos = "l"
        grafico.x_axis.tickLblPos = "low"

    if inteiro:
        grafico.y_axis.number_format = "0"
    _configurar_escala(grafico.y_axis, maximo, divisoes=8, inteiro=inteiro)

    _preparar_grafico(grafico)
    ws.add_chart(grafico, posicao)


def _gravar_tabela_auxiliar(dashboard, produtos, faturamento, quantidades, datas):
    """Escreve a tabela que serve de origem aos gráficos."""
    inicio = COLUNA_AUXILIAR
    titulos = ["Produto", "Faturamento", "Quantidade", "Data", "Faturamento diário"]
    borda = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))

    for deslocamento, titulo in enumerate(titulos):
        celula = dashboard.cell(1, inicio + deslocamento, titulo)
        celula.font = Font(name="Calibri", size=11, bold=True)
        celula.fill = PatternFill("solid", fgColor="E2F0D9")

    for linha, produto in enumerate(produtos, start=2):
        dashboard.cell(linha, inicio, produto)
        dashboard.cell(linha, inicio + 1, faturamento.get(produto, 0)).number_format = FORMATO_MOEDA
        dashboard.cell(linha, inicio + 2, quantidades.get(produto, 0))

    for linha, dia in enumerate(sorted(datas), start=2):
        dashboard.cell(linha, inicio + 3, dia).number_format = FORMATO_DATA
        dashboard.cell(linha, inicio + 4, datas[dia]).number_format = FORMATO_MOEDA

    for linha in dashboard.iter_rows(min_row=1, max_row=max(len(produtos), len(datas)) + 1, min_col=inicio, max_col=inicio + 4):
        for celula in linha:
            celula.border = borda
            celula.alignment = Alignment(horizontal="center", vertical="center")


def _criar_grafico_pizza(dashboard, ultima_linha) -> None:
    inicio = COLUNA_AUXILIAR
    grafico = PieChart()
    grafico.title = "Participação no faturamento"
    grafico.height = 8
    grafico.width = 13

    dados = Reference(dashboard, min_col=inicio + 1, min_row=1, max_row=ultima_linha)
    categorias = Reference(dashboard, min_col=inicio, min_row=2, max_row=ultima_linha)
    grafico.add_data(dados, titles_from_data=True)
    grafico.set_categories(categorias)

    rotulos = DataLabelList()
    rotulos.showPercent = True
    rotulos.showVal = False
    rotulos.showCatName = False
    rotulos.showSerName = False
    rotulos.showLegendKey = False
    rotulos.showLeaderLines = True
    grafico.dataLabels = rotulos

    _preparar_grafico(grafico)
    dashboard.add_chart(grafico, "I3")


def _criar_grafico_linha(dashboard, ultima_linha, maximo=None) -> None:
    inicio = COLUNA_AUXILIAR
    grafico = LineChart()
    grafico.title = "Faturamento por data"
    grafico.style = 13
    grafico.y_axis.title = "Faturamento (R$)"
    grafico.x_axis.title = "Data"
    grafico.height = 9
    grafico.width = 17

    dados = Reference(dashboard, min_col=inicio + 4, min_row=1, max_row=ultima_linha)
    categorias = Reference(dashboard, min_col=inicio + 3, min_row=2, max_row=ultima_linha)
    grafico.add_data(dados, titles_from_data=True)
    grafico.set_categories(categorias)
    grafico.legend = None

    # marcadores.
    serie = grafico.series[0]
    serie.smooth = False
    serie.marker = Marker(symbol="circle", size=6)

    # Posição dos eixos: datas embaixo, valores à esquerda.
    grafico.x_axis.axPos = "b"
    grafico.y_axis.axPos = "l"
    grafico.x_axis.tickLblPos = "low"

    # Um rótulo de data a cada N pontos, para não sobrepor.
    salto = max(1, (ultima_linha - 1 + 4) // 5)
    grafico.x_axis.tickLblSkip = salto
    grafico.x_axis.tickMarkSkip = salto

    # Escala do eixo vertical a partir de zero, com valores sem centavos.
    grafico.y_axis.number_format = '"R$" #,##0'
    _configurar_escala(grafico.y_axis, maximo, divisoes=5)

    _preparar_grafico(grafico)
    dashboard.add_chart(grafico, "A23")


def _criar_dashboard(wb, ws_vendas) -> None:
    """Cria a aba Dashboard com quatro gráficos."""
    if NOME_DASHBOARD in wb.sheetnames:
        del wb[NOME_DASHBOARD]

    dashboard = wb.create_sheet(NOME_DASHBOARD)
    dashboard.sheet_view.showGridLines = False

    dashboard.merge_cells("A1:N1")
    titulo = dashboard["A1"]
    titulo.value = "DASHBOARD DE VENDAS"
    titulo.font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    titulo.fill = PatternFill("solid", fgColor=COR_CABECALHO)
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    dashboard.row_dimensions[1].height = 35

    faturamento, quantidades, datas = _agrupar_vendas(ws_vendas)

    if not faturamento and not quantidades and not datas:
        dashboard["A3"] = "Ainda não existem vendas para gerar os gráficos."
        dashboard["A3"].font = Font(
            name="Calibri", size=12, italic=True, color=COR_CABECALHO
        )
        dashboard.column_dimensions["A"].width = 55
        return

    produtos = sorted(
        set(faturamento) | set(quantidades),
        key=lambda produto: faturamento.get(produto, 0),
        reverse=True,
    )
    _gravar_tabela_auxiliar(dashboard, produtos, faturamento, quantidades, datas)

    ultima_linha_produtos = len(produtos) + 1
    ultima_linha_datas = len(datas) + 1

    if any(valor > 0 for valor in faturamento.values()):
        _criar_grafico_barras(
            dashboard, "Faturamento por produto",
            COLUNA_AUXILIAR + 1, COLUNA_AUXILIAR, ultima_linha_produtos, "A3",
            tipo="bar", titulo_eixo="Faturamento (R$)",
        )
        _criar_grafico_pizza(dashboard, ultima_linha_produtos)

    if datas:
        _criar_grafico_linha(dashboard, ultima_linha_datas, max(datas.values()))

    if any(valor > 0 for valor in quantidades.values()):
        _criar_grafico_barras(
            dashboard, "Quantidade vendida por produto",
            COLUNA_AUXILIAR + 2, COLUNA_AUXILIAR, ultima_linha_produtos, "I23",
            tipo="bar", titulo_eixo="Unidades",
            maximo=max(quantidades.values()), inteiro=True, altura=11,
        )

    for coluna in range(1, 15):
        dashboard.column_dimensions[get_column_letter(coluna)].width = 13
    # A tabela auxiliar fica visível (colunas T a X), à direita dos gráficos.
    for coluna in range(COLUNA_AUXILIAR, COLUNA_AUXILIAR + 5):
        dashboard.column_dimensions[get_column_letter(coluna)].width = 20


def formatar_planilha(caminho: Path) -> None:
    """Formata Vendas.xlsx e cria a aba Dashboard."""
    wb = None

    try:
        wb = load_workbook(caminho)

        if NOME_ABA not in wb.sheetnames:
            raise ErroDeFormatacao(
                f"A aba '{NOME_ABA}' não foi encontrada em '{Path(caminho).name}'."
            )

        ws = wb[NOME_ABA]

        lado = Side(style="medium", color=COR_BORDA)
        borda = Border(left=lado, right=lado, top=lado, bottom=lado)

        _formatar_cabecalho(ws, borda)
        _formatar_corpo(ws, borda)
        _ajustar_larguras(ws)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = (f"A1:{get_column_letter(max(ws.max_column, 1))}{max(ws.max_row, 1)}")

        _criar_dashboard(wb, ws)

        wb.save(caminho)
        print(f"Bordas aplicadas e arquivo salvo: {caminho.name}")

    except ErroDeFormatacao:
        raise
    except PermissionError as erro:
        raise ErroDeFormatacao(
            f"Sem permissão para formatar '{Path(caminho).name}'. "
            "Feche o arquivo no Excel e tente novamente."
        ) from erro
    except (BadZipFile, InvalidFileException) as erro:
        raise ErroDeFormatacao(
            f"'{Path(caminho).name}' não pôde ser aberto para formatação: {erro}"
        ) from erro
    except OSError as erro:
        raise ErroDeFormatacao(
            f"Falha ao formatar '{Path(caminho).name}': {erro}"
        ) from erro
    finally:
        if wb is not None:
            wb.close()
