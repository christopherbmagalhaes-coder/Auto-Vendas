import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def processar_automacao_vendas():
  # 1. Leitura das planilhas de origem
  # (Certifique-se de que os arquivos/abas estejam acessíveis no diretório)
  df_clientes = pd.read_excel('Clientes.xlsx')
  df_estoque = pd.read_excel('Estoque.xlsx')
  df_pedidos = pd.read_excel('Pedidos.xlsx')

  # 2. Cruzamento dos dados (Merge) utilizando o pandas
  # Relacionando Pedidos com Clientes através de 'ID Cliente'
  df_vendas = pd.merge(df_pedidos, df_clientes, on='ID Cliente', how='left')

  # Relacionando com Estoque através de 'ID Produto' para trazer o nome do produto e preço
  df_vendas = pd.merge(df_vendas, df_estoque, on='ID Produto', how='left')

  # 3. Tratamento e organização das colunas finais
  # Gerando um ID sequencial para cada registro da tabela final
  df_vendas['ID'] = range(1, len(df_vendas) + 1)

  # Garantindo que o cálculo de 'Total' esteja correto (Qtd. * Preço unit.)
  # Caso o arquivo de pedidos já traga o total pronto, essa linha pode ser ajustada
  df_vendas['Total'] = df_vendas['Qtd.'] * df_vendas['Preço unit.']

  # Selecionando e reordenando as 8 colunas exigidas para a planilha Vendas
  df_final = df_vendas[[
      'ID',
      'ID Pedido',  # Mapeado para a coluna 'Pedido'
      'Data',
      'Cliente',
      'Produto',
      'Qtd.',
      'Preço unit.',
      'Total',
  ]]

  # Renomeando 'ID Pedido' para 'Pedido' conforme especificado
  df_final = df_final.rename(columns={'ID Pedido': 'Pedido'})

  # 4. Salvando o resultado preliminar em Excel
  arquivo_saida = 'Vendas.xlsx'
  df_final.to_excel(arquivo_saida, index=False)

  # 5. Formatação e Acabamento com openpyxl
  wb = openpyxl.load_workbook(arquivo_saida)
  ws = wb.active
  ws.title = 'Vendas'

  # Estilos visuais profissionais
  font_cabecalho = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
  fill_cabecalho = PatternFill(
      start_color='1F4E78', end_color='1F4E78', fill_type='solid'
  )
  borda_fina = Border(
      left=Side(style='thin', color='D9D9D9'),
      right=Side(style='thin', color='D9D9D9'),
      top=Side(style='thin', color='D9D9D9'),
      bottom=Side(style='thin', color='D9D9D9'),
  )
  alinhamento_central = Alignment(horizontal='center', vertical='center')
  alinhamento_direita = Alignment(horizontal='right', vertical='center')
  alinhamento_esquerda = Alignment(horizontal='left', vertical='center')

  # Aplicando formatação no cabeçalho
  for col in range(1, ws.max_column + 1):
    celula = ws.cell(row=1, column=col)
    celula.font = font_cabecalho
    celula.fill = fill_cabecalho
    celula.alignment = alinhamento_central

  # Aplicando formatação nas linhas de dados
  for row in range(2, ws.max_row + 1):
    for col in range(1, ws.max_column + 1):
      celula = ws.cell(row=row, column=col)
      celula.border = borda_fina

      # Alinhamentos específicos por tipo de dado/coluna
      if col in [1, 2, 3]:  # ID, Pedido, Data
        celula.alignment = alinhamento_central
      elif col in [6]:  # Qtd.
        celula.alignment = alinhamento_right = alinhamento_direita
        celula.number_format = '#,##0'
      elif col in [7, 8]:  # Preço unit. e Total
        celula.alignment = alinhamento_direita
        celula.number_format = 'R$ #,##0.00'
      else:
        celula.alignment = alinhamento_esquerda

  # Ajuste automático da largura das colunas
  for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for celula in col:
      if celula.value:
        max_len = max(max_len, len(str(celula.value)))
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

  # Salvando o arquivo formatado
  wb.save(arquivo_saida)
  print(f"Planilha '{arquivo_saida}' gerada e formatada com sucesso!")


if __name__ == '__main__':
  processar_automacao_vendas()